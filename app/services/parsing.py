"""
Парсинг файлов -> "сырой" текст + разметка структуры.

extract_text_with_pages() возвращает чистый текст документа (без каких-либо
служебных маркеров — это ровно то, что видит пользователь в
GET /documents/{id}/text) вместе со списком PageBreak: где в этом тексте
начинается каждая страница PDF. chunking.py использует page_breaks, чтобы
проставить page_number чанку по его смещению в тексте, вместо старого
подхода со вставкой "--PAGE:N--" прямо в текст (тот подход утекал наружу —
маркер первой страницы не отрезался strip()'ом текста и был виден
пользователю).

Заголовки: "# Текст" — псевдо-markdown заголовок. Для DOCX формируется из
реального стиля параграфа (Heading N / Заголовок N), для PDF/TXT заголовки
определяются эвристикой уже в chunking.py (там текста со стилями нет).
"""

import re
from dataclasses import dataclass
from pathlib import Path
import pymupdf
import docx
from app.models.documents import DocumentFileType

_TXT_ENCODINGS = ("utf-8", "utf-8-sig", "cp1251")
_HEADING_LEVEL_RE = re.compile(r"(\d+)")


@dataclass
class PageBreak:
    """offset — позиция в чистом тексте, с которой начинается страница page_number."""
    offset: int
    page_number: int


class ParsingError(Exception):
    pass


def extract_text(storage_path: str, file_type: DocumentFileType) -> str:
    """Текст документа без разметки страниц — используется там, где номера
    страниц не нужны (например, GET /documents/{id}/text)."""
    return extract_text_with_pages(storage_path, file_type)[0]


def extract_text_with_pages(
    storage_path: str, file_type: DocumentFileType
) -> tuple[str, list[PageBreak]]:
    path = Path(storage_path)
    match file_type:
        case DocumentFileType.txt:
            return _extract_txt(path), []
        case DocumentFileType.pdf:
            return _extract_pdf(path)
        case DocumentFileType.docx:
            return _extract_docx(path), []
        case _:
            raise ParsingError(f"Неподдерживаемый тип файла: {file_type}")


def _extract_txt(path: Path) -> str:
    last_error: UnicodeDecodeError | None = None
    for encoding in _TXT_ENCODINGS:
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError as exc:
            last_error = exc
            continue
    # если ни одна известная кодировка не подошла — берём utf-8 с заменой
    # нечитаемых байт, лучше отдать что-то, чем упасть на ровном месте
    try:
        return path.read_text(encoding="utf-8", errors="replace")
    except Exception as exc:
        raise ParsingError(f"Не удалось прочитать TXT-файл: {last_error or exc}") from exc


def _extract_pdf(path: Path) -> tuple[str, list[PageBreak]]:
    try:
        doc = pymupdf.open(str(path))
    except Exception as exc:
        raise ParsingError(f"Ошибка открытия PDF: {exc}") from exc

    try:
        if doc.needs_pass:
            try:
                if not doc.authenticate(""):
                    raise ParsingError("PDF защищён паролем, извлечение невозможно")
            except ParsingError:
                raise
            except Exception as exc:
                raise ParsingError("PDF защищён паролем, извлечение невозможно") from exc

        parts: list[str] = []
        page_breaks: list[PageBreak] = []
        pos = 0
        for i, page in enumerate(doc, start=1):
            try:
                text = page.get_text().strip()
            except Exception as exc:
                raise ParsingError(f"Ошибка чтения страницы {i} PDF: {exc}") from exc

            if parts:
                sep = "\n\n"
                parts.append(sep)
                pos += len(sep)

            page_breaks.append(PageBreak(offset=pos, page_number=i))
            parts.append(text)
            pos += len(text)
    finally:
        doc.close()

    full_text = "".join(parts)
    if not full_text.strip():
        raise ParsingError(
            "Не удалось извлечь текст из PDF (возможно, это скан без OCR-слоя)"
        )
    return full_text, page_breaks


def _heading_prefix(style_name: str) -> str | None:
    """Heading 1 / Заголовок 2 / Title -> '#', '##', '#' ..."""
    lower = style_name.lower()
    if lower == "title":
        return "#"
    if lower.startswith(("heading", "заголовок")):
        level_match = _HEADING_LEVEL_RE.search(lower)
        level = min(int(level_match.group(1)), 6) if level_match else 1
        return "#" * level
    return None


def _extract_docx(path: Path) -> str:
    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise ParsingError(f"Ошибка чтения DOCX: {exc}") from exc

    lines: list[str] = []
    try:
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            prefix = _heading_prefix(paragraph.style.name if paragraph.style else "")
            lines.append(f"{prefix} {text}" if prefix else text)

        # Таблицы тоже несут смысловое содержимое — добавляем построчно
        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells)
                if row_text.strip(" |"):
                    lines.append(row_text)
    except Exception as exc:
        raise ParsingError(f"Ошибка разбора содержимого DOCX: {exc}") from exc

    text = "\n\n".join(lines).strip()
    if not text:
        raise ParsingError("DOCX не содержит извлекаемого текста")
    return text
